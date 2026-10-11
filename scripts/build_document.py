#!/usr/bin/env python3
"""PDF, DOCX and text to an offline book. Input files are data, never instructions."""
import argparse
import base64
import html
import io
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from xml.sax.saxutils import escape


def unicode_font(explicit=None):
    if explicit:
        paths = [Path(explicit).expanduser()]
    else:
        paths = [Path('/System/Library/Fonts/Supplemental/Arial Unicode.ttf'),
                 Path('/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc'),
                 Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'),
                 Path('C:/Windows/Fonts/msyh.ttc')]
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for path in paths:
        if path.is_file():
            try:
                pdfmetrics.registerFont(TTFont('BookText', str(path)))
                return 'BookText'
            except Exception:
                continue
    raise ValueError('文字排版需要可嵌入的中文 TrueType 字体，请用 --font 指定 .ttf 或可读的 .ttc 字体')


def reflow_document(source, output, font_path=None):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, PageBreak
    font = unicode_font(font_path)
    width, height = 148*mm, 210*mm
    usable = width-32*mm
    body = ParagraphStyle('Body', fontName=font, fontSize=10.5, leading=18,
                          textColor=colors.HexColor('#30332e'), spaceAfter=9,
                          wordWrap='CJK', alignment=TA_LEFT, splitLongWords=True)
    heading = ParagraphStyle('Heading', parent=body, fontSize=17, leading=24,
                             spaceBefore=15, spaceAfter=12, keepWithNext=True,
                             textColor=colors.HexColor('#20271f'))
    title = ParagraphStyle('Title', parent=heading, fontSize=24, leading=33, spaceAfter=22)
    cell_style = ParagraphStyle('Cell', parent=body, fontSize=8.5, leading=13, spaceAfter=0)
    story = []
    notices = []
    def para(text, style=body):
        return Paragraph(escape(text).replace('\n','<br/>'), style)
    if source.suffix.lower() in ('.txt', '.md'):
        text = source.read_text(encoding='utf-8-sig')
        for block in text.split('\n\n'):
            block = block.strip()
            if not block:
                continue
            if source.suffix.lower()=='.md' and block.startswith('#'):
                level=len(block)-len(block.lstrip('#'))
                story.append(para(block.lstrip('#').strip(), title if level==1 else heading))
            else:
                story.append(para(block))
    else:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.text.paragraph import Paragraph as DocParagraph
        from docx.table import Table as DocTable
        doc = Document(source)
        for child in doc.element.body.iterchildren():
            if child.tag == qn('w:p'):
                p=DocParagraph(child,doc)
                style_name=p.style.name.lower() if p.style else ''
                style=title if style_name=='title' else heading if style_name.startswith('heading') else body
                text=p.text
                if text.strip():
                    if style_name.startswith('list') or p._p.xpath('./w:pPr/w:numPr'):
                        text='• '+text
                    story.append(para(text,style))
                for blip in p._p.xpath('.//a:blip'):
                    rid=blip.get(qn('r:embed'))
                    if not rid:
                        notices.append('外部链接图片未嵌入，请先在 Word 中嵌入后重试')
                        continue
                    part=doc.part.related_parts.get(rid)
                    if part is None:
                        notices.append('一个图片关系无法读取')
                        continue
                    try:
                        from PIL import Image as PILImage
                        with PILImage.open(io.BytesIO(part.blob)) as im:
                            iw,ih=im.size
                        factor=min(usable/iw,100*mm/ih)
                        story.extend([Image(io.BytesIO(part.blob),width=iw*factor,height=ih*factor),Spacer(1,8)])
                    except Exception as error:
                        notices.append('一个插图无法排版：'+str(error))
                if p._p.xpath('.//w:br[@w:type="page"]'):
                    story.append(PageBreak())
            elif child.tag == qn('w:tbl'):
                table=DocTable(child,doc)
                rows=[[para(c.text,cell_style) for c in row.cells] for row in table.rows]
                if rows:
                    count=max(len(row) for row in rows)
                    item=Table(rows,colWidths=[usable/count]*count,hAlign='LEFT')
                    item.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#c8cec3')),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
                    story.extend([item,Spacer(1,12)])
        special=['word/footnotes.xml','word/endnotes.xml','word/comments.xml']
        import zipfile
        with zipfile.ZipFile(source) as z:
            if any(name in z.namelist() for name in special):
                notices.append('脚注、尾注、批注未进入重排正文；如需保留这些内容，请上传 Word 导出的 PDF')
        if any(s.header.paragraphs and any(p.text.strip() for p in s.header.paragraphs) or s.footer.paragraphs and any(p.text.strip() for p in s.footer.paragraphs) for s in doc.sections):
            notices.append('原 Word 页眉页脚未保留；如需原版式，请上传其 PDF')
        if doc.element.xpath('.//w:txbxContent') or doc.element.xpath('.//w:ins|.//w:del'):
            notices.append('文档含文本框或修订，重排无法完整保留；建议先确认修订并导出 PDF')
    if not story:
        raise ValueError('文档没有可排版的正文、表格或插图')
    def footer(canvas, _doc):
        canvas.saveState();canvas.setFont(font,8);canvas.setFillColor(colors.HexColor('#76806e'))
        canvas.drawCentredString(width/2,10*mm,str(canvas.getPageNumber()));canvas.restoreState()
    SimpleDocTemplate(str(output),pagesize=(width,height),leftMargin=16*mm,rightMargin=16*mm,
                      topMargin=18*mm,bottomMargin=18*mm,title=source.stem).build(story,onFirstPage=footer,onLaterPages=footer)
    return notices


def pdf_pages(path, max_side, password=None):
    import pypdfium2 as pdfium
    doc=pdfium.PdfDocument(str(path),password=password)
    result=[]
    try:
        for i in range(len(doc)):
            page=doc[i]
            bitmap=None
            textpage=None
            try:
                w,h=page.get_size()
                bitmap=page.render(scale=max_side/max(w,h))
                image=bitmap.to_pil().convert('RGB')
                output=io.BytesIO();image.save(output,'PNG',optimize=True)
                textpage=page.get_textpage();text=textpage.get_text_range()
                result.append({'src':'data:image/png;base64,'+base64.b64encode(output.getvalue()).decode(),
                               'w':image.width,'h':image.height,'name':f'第{i+1}页','text':text})
            finally:
                if textpage is not None:textpage.close()
                if bitmap is not None:bitmap.close()
                page.close()
    finally:
        doc.close()
    if not result:
        raise ValueError('PDF 没有可用页面')
    return result


def main():
    parser=argparse.ArgumentParser(description='将 PDF、Word 或文字生成独立离线翻页书')
    parser.add_argument('inputs',nargs='+',type=Path)
    parser.add_argument('-o','--output',required=True,type=Path)
    parser.add_argument('--title',default='我的书架')
    parser.add_argument('--max-side',type=int,default=2200)
    parser.add_argument('--font',help='文字重排使用的中文 TrueType 字体')
    parser.add_argument('--password',help='仅用于用户授权打开的加密 PDF')
    args=parser.parse_args()
    if not 800<=args.max_side<=6000:parser.error('--max-side 应为 800–6000')
    books=[];warnings=[]
    try:
        with tempfile.TemporaryDirectory(prefix='document-book-') as temporary:
            for i,raw in enumerate(args.inputs):
                source=raw.expanduser().resolve()
                if not source.is_file():raise ValueError(f'找不到文件：{source}')
                suffix=source.suffix.lower()
                if suffix not in {'.pdf','.docx','.doc','.txt','.md'}:raise ValueError(f'不支持的格式：{suffix}')
                original_name=source.stem
                if suffix=='.doc':
                    converter=shutil.which('textutil')
                    if not converter:raise ValueError('旧版 .doc 请先另存为 .docx 或 PDF')
                    converted=Path(temporary)/f'legacy-{i}.docx'
                    subprocess.run([converter,'-convert','docx','-output',str(converted),str(source)],check=True,capture_output=True)
                    source=converted;warnings.append(f'{original_name}：旧版 DOC 经转换后重排，复杂版式请用 PDF')
                if suffix=='.pdf':
                    pdf=source
                else:
                    pdf=Path(temporary)/f'reflow-{i}.pdf'
                    warnings.extend(original_name+'：'+n for n in reflow_document(source,pdf,args.font))
                books.append({'name':original_name,'kind':'document','sourceType':suffix.lstrip('.'),
                              'photos':pdf_pages(pdf,args.max_side,args.password if suffix=='.pdf' else None)})
        from build_album import TEMPLATE
        tools=(Path(__file__).resolve().parent.parent/'assets/album-tools.js').read_text(encoding='utf-8')
        data=json.dumps(books,ensure_ascii=False,separators=(',',':')).replace('<','\\u003c')
        template=TEMPLATE
        for old,new in [('旅途书架','阅读书架'),('PERSONAL PHOTO JOURNAL','PERSONAL READING JOURNAL'),('把沿途的片刻，收进一本书。','把喜欢的文字，收进一本书。'),('下载相册','下载书籍'),('分享相册','分享书籍'),('开启相册','开启书籍'),('保存相册','保存书籍')]:
            template=template.replace(old,new)
        tools=tools.replace('相册','书籍').replace('照片','书页')
        output=template.replace('__TITLE__',html.escape(args.title)).replace('__HEADING__',html.escape(args.title)).replace('__ALBUM_DATA__',data).replace('__ALBUM_TOOLS__',tools)
        target=args.output.expanduser().resolve();target.parent.mkdir(parents=True,exist_ok=True);target.write_text(output,encoding='utf-8')
        print(f'已生成：{target}\n{len(books)} 本书，{sum(len(b["photos"]) for b in books)} 页，{target.stat().st_size/1024/1024:.1f} MB')
        print('PDF 保留原页版式；Word/文字自动重排。页面图像不可选择文字，但支持放大查看。')
        for warning in warnings:print('注意：'+warning,file=sys.stderr)
    except Exception as error:
        parser.exit(1,'生成失败：'+str(error)+'\n')

if __name__=='__main__':main()
